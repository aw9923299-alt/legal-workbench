from __future__ import annotations

import argparse
import json
import os
import re
import subprocess
import sys
import tempfile
from collections.abc import Mapping, Sequence
from dataclasses import asdict, dataclass
from hashlib import sha256
from pathlib import Path
from typing import Any
from zipfile import BadZipFile, ZipFile

from docx import Document
from pypdf import PdfReader

from legal_workbench.domain.entities import ExtractedDocument, ExtractedSegment
from legal_workbench.domain.enums import DocumentExtractionStatus
from legal_workbench.domain.errors import StorageQuotaExceededError

EXTRACTOR_VERSION = "document-text-v1"
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
DOCM_MIME = "application/vnd.ms-word.document.macroenabled.12"
MARKDOWN_MIMES = {"text/markdown", "text/x-markdown"}


class DocumentExtractionError(RuntimeError):
    code = "DOCUMENT_EXTRACTION_FAILED"


class UnsafeAttachmentPathError(DocumentExtractionError):
    code = "UNSAFE_ATTACHMENT_PATH"


class AttachmentTooLargeError(DocumentExtractionError):
    code = "ATTACHMENT_TOO_LARGE"


class ExtractionProcessTimeoutError(DocumentExtractionError):
    code = "DOCUMENT_EXTRACTION_TIMEOUT"


class ExtractionOutputTooLargeError(DocumentExtractionError):
    code = "DOCUMENT_EXTRACTION_OUTPUT_TOO_LARGE"


class ExtractionProcessError(DocumentExtractionError):
    code = "DOCUMENT_EXTRACTION_PROCESS_FAILED"


@dataclass(frozen=True, slots=True)
class AttachmentPathPolicy:
    root: Path
    max_file_bytes: int

    def authorize(self, path: Path) -> Path:
        root = self.root.resolve(strict=True)
        try:
            candidate = path.resolve(strict=True)
            candidate.relative_to(root)
        except (FileNotFoundError, ValueError) as exc:
            raise UnsafeAttachmentPathError(
                "Attachment path is outside the authorized root."
            ) from exc
        if not candidate.is_file():
            raise UnsafeAttachmentPathError("Authorized attachment path is not a file.")
        if candidate.stat().st_size > self.max_file_bytes:
            raise AttachmentTooLargeError("Attachment exceeds the configured parsing limit.")
        return candidate


@dataclass(frozen=True, slots=True)
class StorageQuotaPolicy:
    total_bytes: int

    def ensure_available(
        self,
        *,
        used_bytes: int,
        reserved_bytes: int,
        incoming_bytes: int,
    ) -> None:
        if min(used_bytes, reserved_bytes, incoming_bytes) < 0:
            raise ValueError("Storage quota byte counts cannot be negative.")
        if used_bytes + reserved_bytes + incoming_bytes > self.total_bytes:
            raise StorageQuotaExceededError(
                "Attachment storage quota does not have enough free bytes."
            )


def sanitize_attachment_filename(value: str) -> str:
    base = Path(value.replace("\\", "/")).name.strip() or "attachment"
    sanitized = re.sub(r"[^\w.()\-\u4e00-\u9fff]+", "_", base)
    sanitized = sanitized.lstrip(".")[:240]
    return sanitized or "attachment"


def safe_extraction_environment(
    source: Mapping[str, str] | None = None,
) -> dict[str, str]:
    environment = source if source is not None else os.environ
    allowed = ("LANG", "LC_ALL", "PATH", "SYSTEMROOT", "TMPDIR")
    return {
        key: environment[key] for key in allowed if key in environment and environment[key].strip()
    }


def extract_document(path: Path, mime_type: str) -> ExtractedDocument:
    suffix = path.suffix.lower()
    normalized_mime = mime_type.strip().lower()
    if normalized_mime == "application/pdf" and suffix == ".pdf":
        return _extract_pdf(path)
    if normalized_mime == DOCX_MIME and suffix == ".docx":
        return _extract_docx(path)
    if normalized_mime == "text/plain" and suffix == ".txt":
        return _extract_text(path)
    if normalized_mime in MARKDOWN_MIMES and suffix in {".md", ".markdown"}:
        return _extract_text(path)
    return ExtractedDocument(
        status=DocumentExtractionStatus.BODY_UNAVAILABLE,
        segments=(),
        page_count=None,
        character_count=0,
        error_code="UNSUPPORTED_DOCUMENT_FORMAT",
    )


def _extract_text(path: Path) -> ExtractedDocument:
    try:
        text = path.read_text(encoding="utf-8")
    except UnicodeDecodeError:
        return ExtractedDocument(
            status=DocumentExtractionStatus.FAILED,
            segments=(),
            page_count=None,
            character_count=0,
            error_code="DOCUMENT_TEXT_NOT_UTF8",
        )
    segments = _segments_from_text(text, page_number=None, paragraph_start=1, offset_start=0)
    return _result_from_segments(segments, page_count=None)


def _extract_docx(path: Path) -> ExtractedDocument:
    try:
        with ZipFile(path) as package:
            members = package.infolist()
            if any(
                value.filename.lower().endswith("vbaproject.bin") or value.flag_bits & 0x1
                for value in members
            ):
                return _unsupported_document()
            if sum(value.file_size for value in members) > 100 * 1024 * 1024:
                raise AttachmentTooLargeError(
                    "DOCX expanded content exceeds the safe parsing limit."
                )
        document = Document(str(path))
    except (BadZipFile, KeyError, ValueError) as exc:
        raise DocumentExtractionError("DOCX package is invalid.") from exc
    segments: list[ExtractedSegment] = []
    offset = 0
    for paragraph_number, paragraph in enumerate(document.paragraphs, start=1):
        content = paragraph.text.strip()
        if not content:
            continue
        segments.append(
            _segment(
                content,
                page_number=None,
                paragraph_number=paragraph_number,
                start_offset=offset,
            )
        )
        offset += len(content) + 2
    return _result_from_segments(tuple(segments), page_count=None)


def _extract_pdf(path: Path) -> ExtractedDocument:
    reader = PdfReader(str(path), strict=False)
    segments: list[ExtractedSegment] = []
    offset = 0
    paragraph_number = 1
    for page_number, page in enumerate(reader.pages, start=1):
        page_text = page.extract_text() or ""
        page_segments = _segments_from_text(
            page_text,
            page_number=page_number,
            paragraph_start=paragraph_number,
            offset_start=offset,
        )
        segments.extend(page_segments)
        if page_segments:
            paragraph_number = page_segments[-1].paragraph_number + 1
            offset = page_segments[-1].end_offset + 2
    return _result_from_segments(tuple(segments), page_count=len(reader.pages))


def _segments_from_text(
    text: str,
    *,
    page_number: int | None,
    paragraph_start: int,
    offset_start: int,
) -> tuple[ExtractedSegment, ...]:
    segments: list[ExtractedSegment] = []
    search_from = 0
    paragraph_number = paragraph_start
    for line in text.splitlines():
        content = line.strip()
        if not content:
            continue
        relative_start = text.find(content, search_from)
        if relative_start < 0:
            relative_start = search_from
        start = offset_start + relative_start
        segments.append(
            _segment(
                content,
                page_number=page_number,
                paragraph_number=paragraph_number,
                start_offset=start,
            )
        )
        paragraph_number += 1
        search_from = relative_start + len(content)
    return tuple(segments)


def _segment(
    content: str,
    *,
    page_number: int | None,
    paragraph_number: int,
    start_offset: int,
) -> ExtractedSegment:
    return ExtractedSegment(
        page_number=page_number,
        paragraph_number=paragraph_number,
        start_offset=start_offset,
        end_offset=start_offset + len(content),
        content=content,
        content_hash=sha256(content.encode("utf-8")).hexdigest(),
    )


def _result_from_segments(
    segments: tuple[ExtractedSegment, ...], *, page_count: int | None
) -> ExtractedDocument:
    if not segments:
        return ExtractedDocument(
            status=DocumentExtractionStatus.BODY_UNAVAILABLE,
            segments=(),
            page_count=page_count,
            character_count=0,
            error_code="DOCUMENT_BODY_UNAVAILABLE",
        )
    return ExtractedDocument(
        status=DocumentExtractionStatus.SUCCEEDED,
        segments=segments,
        page_count=page_count,
        character_count=sum(len(value.content) for value in segments),
    )


def _unsupported_document() -> ExtractedDocument:
    return ExtractedDocument(
        status=DocumentExtractionStatus.BODY_UNAVAILABLE,
        segments=(),
        page_count=None,
        character_count=0,
        error_code="UNSUPPORTED_DOCUMENT_FORMAT",
    )


def _optional_int(value: object) -> int | None:
    if value is None:
        return None
    if isinstance(value, bool) or not isinstance(value, (int, str)):
        raise ExtractionProcessError("Extraction result contains an invalid integer.")
    return int(value)


class IsolatedExtractionProcessRunner:
    def __init__(
        self,
        *,
        attachment_root: Path,
        work_root: Path,
        max_file_bytes: int,
        timeout_seconds: float,
        max_output_bytes: int,
        command: Sequence[str] | None = None,
    ) -> None:
        self._path_policy = AttachmentPathPolicy(attachment_root, max_file_bytes)
        self._attachment_root = attachment_root
        self._work_root = work_root
        self._max_file_bytes = max_file_bytes
        self._timeout_seconds = timeout_seconds
        self._max_output_bytes = max_output_bytes
        self._command = (
            list(command)
            if command is not None
            else [
                sys.executable,
                "-I",
                "-m",
                "legal_workbench.integrations.document_extractors",
            ]
        )

    def extract(self, path: Path, mime_type: str) -> ExtractedDocument:
        authorized = self._path_policy.authorize(path)
        self._work_root.mkdir(parents=True, exist_ok=True, mode=0o700)
        work_directory = Path(tempfile.mkdtemp(prefix="extraction-", dir=self._work_root))
        work_directory.chmod(0o700)
        output_path = work_directory / "result.json"
        command = [
            *self._command,
            "--input",
            str(authorized),
            "--mime-type",
            mime_type,
            "--attachment-root",
            str(self._attachment_root.resolve(strict=True)),
            "--max-file-bytes",
            str(self._max_file_bytes),
            "--output",
            str(output_path),
        ]
        try:
            completed = subprocess.run(
                command,
                cwd=work_directory,
                env=safe_extraction_environment(),
                stdin=subprocess.DEVNULL,
                capture_output=True,
                timeout=self._timeout_seconds,
                check=False,
            )
        except subprocess.TimeoutExpired as exc:
            raise ExtractionProcessTimeoutError(
                "Document extraction process exceeded its time limit."
            ) from exc
        if (
            len(completed.stdout) > self._max_output_bytes
            or len(completed.stderr) > self._max_output_bytes
        ):
            raise ExtractionOutputTooLargeError(
                "Document extraction process output exceeded its limit."
            )
        if completed.returncode != 0:
            raise ExtractionProcessError(
                f"Document extraction process exited with code {completed.returncode}."
            )
        if not output_path.is_file():
            raise ExtractionProcessError("Document extraction process produced no result.")
        if output_path.stat().st_size > self._max_output_bytes:
            raise ExtractionOutputTooLargeError("Document extraction result exceeded its limit.")
        try:
            payload = json.loads(output_path.read_text(encoding="utf-8"))
        except (UnicodeDecodeError, json.JSONDecodeError) as exc:
            raise ExtractionProcessError("Document extraction result is invalid JSON.") from exc
        if not isinstance(payload, dict):
            raise ExtractionProcessError("Document extraction result must be an object.")
        return _document_from_payload(payload)


def _document_to_payload(document: ExtractedDocument) -> dict[str, object]:
    return {
        "status": document.status.value,
        "segments": [asdict(value) for value in document.segments],
        "page_count": document.page_count,
        "character_count": document.character_count,
        "extractor_version": document.extractor_version,
        "error_code": document.error_code,
    }


def _document_from_payload(payload: Mapping[str, Any]) -> ExtractedDocument:
    raw_segments = payload.get("segments")
    if not isinstance(raw_segments, list):
        raise ExtractionProcessError("Extraction result omitted its segment list.")
    segments = tuple(
        ExtractedSegment(
            page_number=_optional_int(value.get("page_number")),
            paragraph_number=int(value["paragraph_number"]),
            start_offset=int(value["start_offset"]),
            end_offset=int(value["end_offset"]),
            content=str(value["content"]),
            content_hash=str(value["content_hash"]),
        )
        for value in raw_segments
        if isinstance(value, dict)
    )
    return ExtractedDocument(
        status=DocumentExtractionStatus(str(payload["status"])),
        segments=segments,
        page_count=_optional_int(payload.get("page_count")),
        character_count=int(payload["character_count"]),
        extractor_version=str(payload["extractor_version"]),
        error_code=(str(payload["error_code"]) if payload.get("error_code") else None),
    )


def _write_private_json(path: Path, payload: dict[str, object]) -> None:
    descriptor = os.open(path, os.O_WRONLY | os.O_CREAT | os.O_EXCL, 0o600)
    with os.fdopen(descriptor, "w", encoding="utf-8") as stream:
        json.dump(payload, stream, ensure_ascii=False, separators=(",", ":"))


def _cli(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", type=Path, required=True)
    parser.add_argument("--mime-type", required=True)
    parser.add_argument("--attachment-root", type=Path, required=True)
    parser.add_argument("--max-file-bytes", type=int, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args(argv)
    authorized = AttachmentPathPolicy(args.attachment_root, args.max_file_bytes).authorize(
        args.input
    )
    result = extract_document(authorized, args.mime_type)
    _write_private_json(args.output, _document_to_payload(result))
    return 0


if __name__ == "__main__":
    raise SystemExit(_cli())
